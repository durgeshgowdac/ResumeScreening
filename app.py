import streamlit as st
import pandas as pd
import joblib
from sklearn.metrics.pairwise import cosine_similarity
from sklearn.feature_extraction.text import TfidfVectorizer
import plotly.express as px
import plotly.graph_objects as go
import docx
import pypdf
import os
from typing import List, Dict, Tuple, Optional
import numpy as np
from datetime import datetime

import warnings
from sklearn.exceptions import InconsistentVersionWarning

warnings.filterwarnings("ignore", category=InconsistentVersionWarning)

import logging
logging.getLogger("pypdf").setLevel(logging.ERROR)

# Page configuration
st.set_page_config(
    page_title="AI Resume Screening System",
    layout="wide",
    page_icon="🧠",
    initial_sidebar_state="expanded"
)

# Custom CSS for better styling
st.markdown("""
<style>
    .metric-card {
        background-color: #f0f2f6;
        padding: 1rem;
        border-radius: 0.5rem;
        border-left: 4px solid #1f77b4;
    }
    .pass-status {
        background-color: #d4edda;
        color: #155724;
        padding: 0.25rem 0.5rem;
        border-radius: 0.25rem;
        font-weight: bold;
    }
    .fail-status {
        background-color: #f8d7da;
        color: #721c24;
        padding: 0.25rem 0.5rem;
        border-radius: 0.25rem;
        font-weight: bold;
    }
    .stExpander > div:first-child {
        background-color: #f8f9fa;
    }
</style>
""", unsafe_allow_html=True)


# =====================
# Configuration Class
# =====================
class Config:
    """Application configuration"""
    DEFAULT_FILE_PATHS = {
        "tfidf_vectorizer": "models/tfidf_vectorizer.pkl",
        "resume_data": "data/processed/resume_data_with_features.parquet",
        "classifier_model": "models/best_model_random_forest.pkl",
        "label_encoder": "models/label_encoder.pkl"
    }

    SUPPORTED_FILE_TYPES = {
        "application/pdf": "PDF",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "DOCX",
        "text/plain": "TXT"
    }

    MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB
    DEFAULT_TFIDF_FEATURES = 5000
    MIN_TEXT_LENGTH = 50


# =====================
# Helper Classes
# =====================
class FileProcessor:
    """Handles file processing operations"""

    @staticmethod
    def extract_text_from_pdf(file) -> str:
        """Extract text from PDF file with better error handling"""
        try:
            pdf_reader = pypdf.PdfReader(file)
            text_parts = []

            for page_num, page in enumerate(pdf_reader.pages):
                try:
                    page_text = page.extract_text()
                    if page_text.strip():
                        text_parts.append(page_text)
                except Exception as e:
                    print(f"Error extracting text from page {page_num}: {e}")
                    continue

            full_text = "\n".join(text_parts)
            return full_text if len(full_text.strip()) >= Config.MIN_TEXT_LENGTH else ""

        except Exception as e:
            print(f"Error reading PDF: {e}")
            return ""

    @staticmethod
    def extract_text_from_docx(file) -> str:
        """Extract text from DOCX file with better error handling"""
        try:
            doc = docx.Document(file)
            text_parts = []

            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text.strip())

            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_parts.append(cell.text.strip())

            full_text = "\n".join(text_parts)
            return full_text if len(full_text.strip()) >= Config.MIN_TEXT_LENGTH else ""

        except Exception as e:
            print(f"Error reading DOCX: {e}")
            return ""

    @staticmethod
    def extract_text_from_txt(file) -> str:
        """Extract text from TXT file with encoding detection"""
        try:
            # Try different encodings
            encodings = ['utf-8', 'latin-1', 'cp1252', 'iso-8859-1']

            for encoding in encodings:
                try:
                    file.seek(0)  # Reset file pointer
                    text = file.read().decode(encoding)
                    return text if len(text.strip()) >= Config.MIN_TEXT_LENGTH else ""
                except UnicodeDecodeError:
                    continue

            # If all encodings fail, try with error handling
            file.seek(0)
            text = file.read().decode('utf-8', errors='ignore')
            return text if len(text.strip()) >= Config.MIN_TEXT_LENGTH else ""

        except Exception as e:
            print(f"Error reading TXT: {e}")
            return ""


class ResumeProcessor:
    """Handles resume processing and matching"""

    def __init__(self):
        self.tfidf_vectorizer = None
        self.classifier = None
        self.label_encoder = None

    def process_uploaded_files(self, uploaded_files) -> pd.DataFrame:
        """Process uploaded resume files with progress tracking"""
        resume_data = []

        if not uploaded_files:
            return pd.DataFrame()

        progress_bar = st.progress(0)
        status_text = st.empty()

        for i, file in enumerate(uploaded_files):
            try:
                # Update progress
                progress = (i + 1) / len(uploaded_files)
                progress_bar.progress(progress)
                status_text.text(f"Processing {file.name}... ({i + 1}/{len(uploaded_files)})")

                # Check file size
                if file.size > Config.MAX_FILE_SIZE:
                    st.warning(f"⚠️ {file.name} exceeds maximum file size (10MB)")
                    continue

                # Extract text based on file type
                text = self._extract_text_from_file(file)

                if text and len(text.strip()) >= Config.MIN_TEXT_LENGTH:
                    resume_data.append({
                        "filename": file.name,
                        "text": text,
                        "category": "Uploaded",
                        "file_size": file.size,
                        "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    })
                else:
                    st.warning(f"⚠️ {file.name}: Insufficient text content")

            except Exception as e:
                print(f"Error processing {file.name}: {e}")
                st.error(f"❌ Error processing {file.name}: {str(e)}")

        # Clear progress indicators
        progress_bar.empty()
        status_text.empty()

        return pd.DataFrame(resume_data)

    def _extract_text_from_file(self, file) -> str:
        """Extract text from file based on type"""
        file_type = file.type

        if file_type == "application/pdf":
            return FileProcessor.extract_text_from_pdf(file)
        elif file_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return FileProcessor.extract_text_from_docx(file)
        elif file_type == "text/plain":
            return FileProcessor.extract_text_from_txt(file)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

    def calculate_similarity_scores(self, job_description: str, resume_texts: List[str]) -> np.ndarray:
        if self.tfidf_vectorizer is None:
            # Disable the ML classifier to prevent dimensional mismatch crashes
            self.classifier = None
            self.label_encoder = None

            # Pull directly from Streamlit session state configured in the sidebar
            max_feat = st.session_state.get('max_features', Config.DEFAULT_TFIDF_FEATURES)
            ngrams = st.session_state.get('ngram_range', (1, 2))

            self.tfidf_vectorizer = TfidfVectorizer(
                max_features=max_feat,
                stop_words='english',
                ngram_range=ngrams,
                min_df=1,
                max_df=0.95
            )
            all_texts = [job_description] + resume_texts
            self.tfidf_vectorizer.fit(all_texts)

        # Transform texts
        jd_vector = self.tfidf_vectorizer.transform([job_description])
        resume_vectors = self.tfidf_vectorizer.transform(resume_texts)

        # Calculate cosine similarity
        similarity_scores = cosine_similarity(jd_vector, resume_vectors).flatten()
        return similarity_scores

    def calculate_hybrid_scores(self, jd_vector, resume_vectors, similarity_scores: np.ndarray,
                                alpha: float, beta: float) -> Tuple[np.ndarray, np.ndarray, str]:
        """Calculate hybrid scores combining similarity and classifier predictions"""
        if self.classifier is None or self.label_encoder is None:
            raise ValueError("Classifier or label encoder not loaded")

        # Predict class for job description
        jd_pred_class = self.classifier.predict(jd_vector)[0]
        jd_class_index = list(self.classifier.classes_).index(jd_pred_class)

        # Get classifier probabilities
        resume_probs = self.classifier.predict_proba(resume_vectors)
        classifier_confidences = resume_probs[:, jd_class_index]

        # Calculate hybrid scores
        hybrid_scores = alpha * similarity_scores + beta * classifier_confidences

        # Get predicted category name
        predicted_category = self.label_encoder.inverse_transform([jd_pred_class])[0]

        return hybrid_scores, classifier_confidences, predicted_category


class AssetLoader:
    """Handles loading of pre-trained models and data"""

    @staticmethod
    @st.cache_resource
    def load_assets(file_paths: Dict[str, str]) -> Tuple:
        """Load pre-trained assets with caching"""
        loading_status = {'tfidf': False, 'data': False, 'classifier': False, 'label_encoder': False}
        try:
            assets = {}
            assets = {}
            loading_status = {}

            # Load TF-IDF vectorizer
            if os.path.exists(file_paths["tfidf_vectorizer"]):
                assets['tfidf'] = joblib.load(file_paths["tfidf_vectorizer"])
                loading_status['tfidf'] = True
            else:
                assets['tfidf'] = None
                loading_status['tfidf'] = False

            # Load resume data
            if os.path.exists(file_paths["resume_data"]):
                df = pd.read_parquet(file_paths["resume_data"])
                df["Text"] = df["Text"].astype(str)
                assets['resume_df'] = df

                if assets['tfidf'] is not None:
                    assets['resume_vectors'] = assets['tfidf'].transform(df["Text"].tolist())
                else:
                    assets['resume_vectors'] = None
                loading_status['data'] = True
            else:
                assets['resume_df'] = None
                assets['resume_vectors'] = None
                loading_status['data'] = False

            # Load classifier
            if os.path.exists(file_paths["classifier_model"]):
                assets['classifier'] = joblib.load(file_paths["classifier_model"])
                loading_status['classifier'] = True
            else:
                assets['classifier'] = None
                loading_status['classifier'] = False

            # Load label encoder
            if os.path.exists(file_paths["label_encoder"]):
                assets['label_encoder'] = joblib.load(file_paths["label_encoder"])
                loading_status['label_encoder'] = True
            else:
                assets['label_encoder'] = None
                loading_status['label_encoder'] = False

            return (assets.get('tfidf'), assets.get('resume_df'), assets.get('resume_vectors'),
                    assets.get('classifier'), assets.get('label_encoder'), loading_status)

        except Exception as e:
            print(f"Error loading assets: {e}")
            return None, None, None, None, None, loading_status


class Visualizer:
    """Handles data visualization"""

    @staticmethod
    def create_pass_fail_pie_chart(status_counts: pd.Series, title: str) -> go.Figure:
        """Create a pie chart for pass/fail distribution"""
        fig = px.pie(
            values=status_counts.values,
            names=status_counts.index,
            title=title,
            color=status_counts.index,
            color_discrete_map={'PASS': '#28a745', 'FAIL': '#dc3545'}
        )
        fig.update_traces(textposition='inside', textinfo='percent+label')
        fig.update_layout(height=400)
        return fig

    @staticmethod
    def create_score_histogram(df: pd.DataFrame, score_column: str, threshold: float, title: str) -> go.Figure:
        """Create histogram with threshold line"""
        fig = px.histogram(
            df,
            x=score_column,
            nbins=20,
            title=title,
            color='status',
            color_discrete_map={'PASS': '#28a745', 'FAIL': '#dc3545'}
        )
        fig.add_vline(
            x=threshold,
            line_dash="dash",
            line_color="black",
            annotation_text=f"Threshold: {threshold:.2f}",
            annotation_position="top"
        )
        fig.update_layout(height=400)
        return fig

    @staticmethod
    def create_hybrid_scatter_plot(df: pd.DataFrame) -> go.Figure:
        """Create scatter plot for hybrid analysis"""
        fig = px.scatter(
            df,
            x="similarity_score",
            y="classifier_confidence",
            size="hybrid_score",
            color="status",
            color_discrete_map={'PASS': '#28a745', 'FAIL': '#dc3545'},
            title="Similarity vs Classifier Confidence",
            hover_data=["hybrid_score", "filename"]
        )
        fig.update_layout(height=500)
        return fig


# =====================
# Main Application
# =====================
def main():
    """Main application function"""

    # Initialize session state
    if 'processor' not in st.session_state:
        st.session_state.processor = ResumeProcessor()

    # Header
    st.title("🧠 AI-Powered Resume Screening System")
    st.markdown("**Advanced resume matching with semantic analysis and pass/fail criteria**")

    # Sidebar configuration
    setup_sidebar()

    # File paths configuration
    file_paths = get_file_paths()

    # Load pre-trained assets
    with st.spinner("Loading pre-trained models..."):
        (tfidf_vectorizer, resume_df, resume_vectors,
         classifier, label_encoder, loading_status) = AssetLoader.load_assets(file_paths)

    # Update processor with loaded assets
    if tfidf_vectorizer:
        st.session_state.processor.tfidf_vectorizer = tfidf_vectorizer
    if classifier:
        st.session_state.processor.classifier = classifier
    if label_encoder:
        st.session_state.processor.label_encoder = label_encoder

    # Display loading status
    display_loading_status(loading_status)

    # Data source selection
    uploaded_resume_df = handle_data_source(resume_df)

    # Job description input
    job_description = handle_job_description_input()

    # Matching configuration
    matching_type, alpha, beta = handle_matching_configuration(classifier, label_encoder)

    # Similarity threshold
    similarity_threshold = st.sidebar.slider(
        "Pass/Fail Threshold (%)",
        0, 100, 50,
        help="Resumes scoring above this threshold will be marked as PASS"
    ) / 100

    # Main matching logic
    if st.button("🔍 Match Resumes", type="primary", width="stretch"):
        is_valid, validation_msg = validate_job_description(job_description)

        if not is_valid:
            st.error(f"❌ {validation_msg}")
        elif uploaded_resume_df is None or uploaded_resume_df.empty:
            st.error("❌ Please upload resumes or select a data source")
        else:
            perform_matching(
                uploaded_resume_df, job_description, matching_type,
                similarity_threshold, alpha, beta
            )

    # Instructions
    display_instructions()


def setup_sidebar():
    """Setup sidebar configuration"""
    st.sidebar.header("⚙️ Configuration")

    # Display options
    st.sidebar.subheader("📊 Display Options")
    st.session_state.num_results = st.sidebar.slider("Number of results to show", 5, 50, 20)
    st.session_state.preview_length = st.sidebar.slider("Resume preview length", 300, 10000, 1500)

    # Performance options
    with st.sidebar.expander("🔧 Advanced Options"):
        st.session_state.max_features = st.slider("TF-IDF max features", 1000, 10000, 5000)
        st.session_state.ngram_range = st.selectbox("N-gram range",
                                                    options=[(1, 1), (1, 2), (1, 3)],
                                                    index=1)


def get_file_paths() -> Dict[str, str]:
    """Get file paths from configuration"""
    # You can modify this to use different path sources
    return Config.DEFAULT_FILE_PATHS


def display_loading_status(loading_status: Dict[str, bool]):
    """Display the status of loaded assets"""
    if loading_status:
        st.sidebar.markdown("### 📁 Asset Status")

        status_icons = {True: "✅", False: "❌"}
        status_labels = {
            'tfidf': "TF-IDF Vectorizer",
            'data': "Resume Dataset",
            'classifier': "ML Classifier",
            'label_encoder': "Label Encoder"
        }

        for key, status in loading_status.items():
            if key in status_labels:
                icon = status_icons[status]
                label = status_labels[key]
                st.sidebar.markdown(f"{icon} {label}")


def handle_data_source(resume_df) -> Optional[pd.DataFrame]:
    """Handle data source selection and processing"""
    st.header("📁 Data Source")

    data_source = st.radio(
        "Choose your data source:",
        ["Upload Resume Files", "Upload CSV File", "Use Pre-loaded Dataset"],
        horizontal=True
    )

    uploaded_resume_df = None

    if data_source == "Upload Resume Files":
        uploaded_resume_df = handle_file_upload()
    elif data_source == "Upload CSV File":
        uploaded_resume_df = handle_csv_upload()
    elif data_source == "Use Pre-loaded Dataset":
        uploaded_resume_df = handle_preloaded_dataset(resume_df)

    return uploaded_resume_df


def handle_file_upload() -> Optional[pd.DataFrame]:
    """Handle individual file uploads"""
    st.subheader("📤 Upload Resume Files")

    uploaded_files = st.file_uploader(
        "Upload resume files (PDF, DOCX, TXT)",
        type=['pdf', 'docx', 'txt'],
        accept_multiple_files=True,
        help="Maximum file size: 10MB per file"
    )

    if uploaded_files:
        uploaded_resume_df = st.session_state.processor.process_uploaded_files(uploaded_files)

        if not uploaded_resume_df.empty:
            st.success(f"✅ Successfully processed {len(uploaded_resume_df)} resumes")

            # Display summary
            col1, col2, col3 = st.columns(3)
            with col1:
                st.metric("Total Files", len(uploaded_resume_df))
            with col2:
                total_size = uploaded_resume_df['file_size'].sum()
                st.metric("Total Size", format_file_size(total_size))
            with col3:
                avg_length = uploaded_resume_df['text'].str.len().mean()
                st.metric("Avg. Text Length", f"{avg_length:.0f} chars")

            # Display file list
            display_df = uploaded_resume_df[['filename', 'category', 'processed_at']].copy()
            st.dataframe(display_df, width="stretch")

            return uploaded_resume_df
        else:
            st.error("❌ No valid resumes found in uploaded files")

    return None


def handle_csv_upload() -> Optional[pd.DataFrame]:
    """Handle CSV file upload"""
    st.subheader("📊 Upload CSV File")
    st.info("CSV should have a 'text' column containing resume content. Optional columns: 'filename', 'category'")

    csv_file = st.file_uploader("Upload CSV file", type=['csv'])

    if csv_file:
        try:
            uploaded_resume_df = pd.read_csv(csv_file)

            # Validate required columns
            if 'text' not in uploaded_resume_df.columns:
                st.error("❌ CSV must contain a 'text' column")
                return None

            # Add missing columns with defaults
            if 'filename' not in uploaded_resume_df.columns:
                uploaded_resume_df['filename'] = [f"Resume_{i + 1}" for i in range(len(uploaded_resume_df))]
            if 'category' not in uploaded_resume_df.columns:
                uploaded_resume_df['category'] = "Uploaded"

            # Filter out empty text entries
            uploaded_resume_df = uploaded_resume_df[uploaded_resume_df['text'].str.len() >= Config.MIN_TEXT_LENGTH]

            if not uploaded_resume_df.empty:
                st.success(f"✅ Successfully loaded {len(uploaded_resume_df)} resumes from CSV")
                st.dataframe(uploaded_resume_df[['filename', 'category']], width='stretch')
                return uploaded_resume_df
            else:
                st.error("❌ No valid resumes found in CSV file")

        except Exception as e:
            st.error(f"❌ Error reading CSV file: {e}")

    return None


def handle_preloaded_dataset(resume_df) -> Optional[pd.DataFrame]:
    """Handle pre-loaded dataset with a fallback to sample data"""
    if resume_df is None:
        st.warning("⚠️ Pre-loaded dataset not found. Loading sample demonstration data instead.")
        # Activate the previously unused utility function
        resume_df = create_sample_data()
        # Align column names to match the expected schema for the rest of the function
        resume_df = resume_df.rename(columns={'text': 'Text', 'category': 'Category'})
    else:
        st.success(f"✅ Using pre-loaded dataset with {len(resume_df)} resumes")

    # Create compatible dataframe
    uploaded_resume_df = resume_df.copy()

    # Use existing filenames from sample data if present, otherwise generate them
    if 'filename' not in uploaded_resume_df.columns:
        uploaded_resume_df['filename'] = [f"Resume_{i + 1}" for i in range(len(uploaded_resume_df))]

    uploaded_resume_df['text'] = uploaded_resume_df['Text']
    uploaded_resume_df['category'] = uploaded_resume_df['Category']

    # Display dataset info
    col1, col2, col3 = st.columns(3)
    with col1:
        st.metric("Total Resumes", len(uploaded_resume_df))
    with col2:
        unique_categories = uploaded_resume_df['category'].nunique()
        st.metric("Categories", unique_categories)
    with col3:
        avg_length = uploaded_resume_df['text'].str.len().mean()
        st.metric("Avg. Text Length", f"{avg_length:.0f} chars")

    return uploaded_resume_df


def handle_job_description_input() -> str:
    """Handle job description input"""
    st.header("📄 Job Description")

    # Option to load from file or paste directly
    input_method = st.radio("Input method:", ["Paste Text", "Upload File"], horizontal=True)

    job_description = ""

    if input_method == "Paste Text":
        job_description = st.text_area(
            "Paste Job Description",
            height=200,
            placeholder="Enter the job description you want to match resumes against..."
        )
    else:
        jd_file = st.file_uploader("Upload JD file", type=['txt', 'pdf', 'docx'])
        if jd_file:
            try:
                if jd_file.type == "text/plain":
                    job_description = FileProcessor.extract_text_from_txt(jd_file)
                elif jd_file.type == "application/pdf":
                    job_description = FileProcessor.extract_text_from_pdf(jd_file)
                elif jd_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                    job_description = FileProcessor.extract_text_from_docx(jd_file)

                if job_description:
                    st.success("✅ Job description loaded from file")
                    st.text_area("Extracted text:", job_description, height=150, disabled=True)

            except Exception as e:
                st.error(f"Error reading file: {e}")

    return job_description


def handle_matching_configuration(classifier, label_encoder) -> Tuple[str, float, float]:
    """Handle matching configuration"""
    st.header("🎯 Matching Configuration")

    # Check if hybrid matching is available
    hybrid_available = classifier is not None and label_encoder is not None

    if hybrid_available:
        matching_options = ["Semantic Similarity", "Hybrid Matching (Similarity + ML Classifier)"]
        st.success("🤖 ML Classifier available - Hybrid matching enabled")
    else:
        matching_options = ["Semantic Similarity"]
        st.info("ℹ️ Using semantic similarity matching (ML classifier not available)")

    matching_type = st.radio("Choose matching method:", matching_options, horizontal=True)

    # Weight configuration for hybrid matching
    alpha, beta = 0.5, 0.5
    if matching_type == "Hybrid Matching (Similarity + ML Classifier)":
        st.subheader("⚖️ Score Balance")

        # A single slider guarantees alpha + beta = 1.0
        alpha = st.slider(
            "Similarity Weight (α)",
            min_value=0.0,
            max_value=1.0,
            value=0.4,
            step=0.1,
            help="Higher values favor semantic similarity. Classifier weight (β) scales automatically."
        )
        beta = round(1.0 - alpha, 1)

        # Display the resulting balance cleanly
        col1, col2 = st.columns(2)
        with col1:
            st.success(f"**Similarity (α):** {alpha:.1f}")
        with col2:
            st.info(f"**Classifier (β):** {beta:.1f}")

    return matching_type, alpha, beta


def perform_matching(uploaded_resume_df: pd.DataFrame, job_description: str,
                     matching_type: str, similarity_threshold: float,
                     alpha: float, beta: float):
    """Perform the actual resume matching"""

    with st.spinner("🔄 Calculating similarities..."):
        processor = st.session_state.processor

        # Calculate similarity scores
        resume_texts = uploaded_resume_df['text'].tolist()
        similarity_scores = processor.calculate_similarity_scores(job_description, resume_texts)

        # Add similarity scores to dataframe
        results_df = uploaded_resume_df.copy()
        results_df['similarity_score'] = similarity_scores

        # Hybrid matching if selected
        if matching_type == "Hybrid Matching (Similarity + ML Classifier)":
            try:
                jd_vector = processor.tfidf_vectorizer.transform([job_description])
                resume_vectors = processor.tfidf_vectorizer.transform(resume_texts)

                hybrid_scores, classifier_confidences, predicted_category = processor.calculate_hybrid_scores(
                    jd_vector, resume_vectors, similarity_scores, alpha, beta
                )

                results_df['classifier_confidence'] = classifier_confidences
                results_df['hybrid_score'] = hybrid_scores
                results_df['status'] = results_df['hybrid_score'].apply(
                    lambda x: "PASS" if x >= similarity_threshold else "FAIL"
                )

                # Sort by hybrid score
                results_df = results_df.sort_values('hybrid_score', ascending=False)

                st.success(f"✅ Hybrid matching complete! Predicted JD Category: **{predicted_category}**")

            except Exception as e:
                st.error(f"❌ Error in hybrid matching: {e}")
                return
        else:
            # Semantic similarity only
            results_df['status'] = results_df['similarity_score'].apply(
                lambda x: "PASS" if x >= similarity_threshold else "FAIL"
            )
            results_df = results_df.sort_values('similarity_score', ascending=False)

    # Display results
    display_results(results_df, matching_type, similarity_threshold, alpha, beta)


def display_results(results_df: pd.DataFrame, matching_type: str,
                    similarity_threshold: float, alpha: float, beta: float):
    """Display matching results with visualizations"""

    # Results summary
    st.header("📊 Results Summary")

    pass_count = (results_df['status'] == 'PASS').sum()
    fail_count = (results_df['status'] == 'FAIL').sum()

    col1, col2, col3, col4 = st.columns(4)
    with col1:
        st.metric("Total Resumes", len(results_df))
    with col2:
        st.metric("✅ Passed", pass_count, delta=f"{pass_count / len(results_df) * 100:.1f}%")
    with col3:
        st.metric("❌ Failed", fail_count, delta=f"{fail_count / len(results_df) * 100:.1f}%")
    with col4:
        st.metric("Threshold", f"{similarity_threshold:.2f}")

    # Visualizations
    st.subheader("📈 Visual Analytics")

    visualizer = Visualizer()
    status_counts = results_df['status'].value_counts()

    if matching_type == "Hybrid Matching (Similarity + ML Classifier)":
        # Hybrid matching visualizations
        col1, col2 = st.columns(2)

        with col1:
            # Pass/Fail pie chart
            fig_pie = visualizer.create_pass_fail_pie_chart(status_counts, "Pass/Fail Distribution (Hybrid)")
            st.plotly_chart(fig_pie, width="stretch")

        with col2:
            # Hybrid score distribution with threshold line
            fig_hist = visualizer.create_score_histogram(
                results_df, 'hybrid_score', similarity_threshold, "Hybrid Score Distribution"
            )
            st.plotly_chart(fig_hist, width='stretch')

        # Additional hybrid visualization - Scatter plot
        st.subheader("🔍 Hybrid Analysis")
        fig_scatter = visualizer.create_hybrid_scatter_plot(results_df.head(st.session_state.num_results))
        st.plotly_chart(fig_scatter, width='stretch')

        # Score component breakdown
        col1, col2, col3 = st.columns(3)
        with col1:
            st.metric("Max Similarity", f"{results_df['similarity_score'].max():.3f}")
        with col2:
            st.metric("Max Classifier Confidence", f"{results_df['classifier_confidence'].max():.3f}")
        with col3:
            st.metric("Max Hybrid Score", f"{results_df['hybrid_score'].max():.3f}")

    else:
        # Semantic similarity visualizations
        col1, col2 = st.columns(2)

        with col1:
            # Pass/Fail pie chart
            fig_pie = visualizer.create_pass_fail_pie_chart(status_counts, "Pass/Fail Distribution (Semantic)")
            st.plotly_chart(fig_pie, width="stretch")

        with col2:
            # Score distribution with threshold line
            fig_hist = visualizer.create_score_histogram(
                results_df, 'similarity_score', similarity_threshold, "Similarity Score Distribution"
            )
            st.plotly_chart(fig_hist, width='stretch')

    # Results table with color coding
    st.header("📋 Results Summary")

    # Display top N results with appropriate columns
    display_df = results_df.head(st.session_state.num_results).copy()

    if matching_type == "Hybrid Matching (Similarity + ML Classifier)":
        # Hybrid matching table
        display_columns = ['filename', 'category', 'similarity_score', 'classifier_confidence', 'hybrid_score',
                           'status']
        display_df_formatted = display_df[display_columns].copy()
        display_df_formatted['similarity_score'] = display_df_formatted['similarity_score'].apply(lambda x: f"{x:.3f}")
        display_df_formatted['classifier_confidence'] = display_df_formatted['classifier_confidence'].apply(
            lambda x: f"{x:.3f}")
        display_df_formatted['hybrid_score'] = display_df_formatted['hybrid_score'].apply(lambda x: f"{x:.3f}")
    else:
        # Semantic similarity table
        display_columns = ['filename', 'category', 'similarity_score', 'status']
        display_df_formatted = display_df[display_columns].copy()
        display_df_formatted['similarity_score'] = display_df_formatted['similarity_score'].apply(lambda x: f"{x:.3f}")

    # Style the dataframe
    def color_status(val):
        color = '#28a745' if val == 'PASS' else '#dc3545'
        return f'background-color: {color}; color: white; font-weight: bold'

    styled_df = display_df_formatted.style.map(color_status, subset=['status'])
    st.dataframe(styled_df, width='stretch')

    # Detailed results with expandable sections
    st.header("📝 Detailed Results")

    # Filter options
    col1, col2 = st.columns(2)
    with col1:
        show_status = st.selectbox("Filter by status:", ["All", "PASS", "FAIL"])
    with col2:
        sort_by = st.selectbox("Sort by:", ["Similarity Score", "Filename", "Status"])

    # Apply filters
    filtered_df = results_df.copy()
    if show_status != "All":
        filtered_df = filtered_df[filtered_df['status'] == show_status]

    # Apply sorting
    if sort_by == "Similarity Score":
        sort_column = 'hybrid_score' if matching_type == "Hybrid Matching (Similarity + ML Classifier)" else 'similarity_score'
        filtered_df = filtered_df.sort_values(sort_column, ascending=False)
    elif sort_by == "Filename":
        filtered_df = filtered_df.sort_values('filename')
    elif sort_by == "Status":
        filtered_df = filtered_df.sort_values('status')

    # Display detailed results
    for i, (_, row) in enumerate(filtered_df.head(st.session_state.num_results).iterrows(), 1):
        status_color = get_status_color(row['status'])

        if matching_type == "Hybrid Matching (Similarity + ML Classifier)":
            # Hybrid matching details
            with st.expander(f"#{i} - {row['filename']} | Hybrid Score: {row['hybrid_score']:.3f} | {row['status']}"):
                col1, col2, col3, col4 = st.columns(4)
                with col1:
                    st.metric("Similarity Score", f"{row['similarity_score']:.3f}")
                with col2:
                    st.metric("Classifier Confidence", f"{row['classifier_confidence']:.3f}")
                with col3:
                    st.metric("Hybrid Score", f"{row['hybrid_score']:.3f}")
                with col4:
                    st.markdown(
                        f"**Status:** <span style='color: {status_color}; font-weight: bold'>{row['status']}</span>",
                        unsafe_allow_html=True)

                # Score breakdown
                st.markdown("### 🧮 Score Components")
                st.markdown(f"**Similarity Weight (α):** {alpha:.1f}")
                st.markdown(f"**Classifier Weight (β):** {beta:.1f}")
                st.markdown(
                    f"**Hybrid Formula:** {alpha:.1f} × {row['similarity_score']:.3f} + {beta:.1f} × {row['classifier_confidence']:.3f} = {row['hybrid_score']:.3f}")

                st.markdown("**Document Metrics:**")
                # Utilize the text statistics utility function
                stats = calculate_text_statistics(row['text'])

                stat_col1, stat_col2, stat_col3 = st.columns(3)
                with stat_col1:
                    st.caption(f"📝 {stats['words']} Words")
                with stat_col2:
                    st.caption(f"📑 {stats['sentences']} Sentences")
                with stat_col3:
                    st.caption(f"📏 {stats['avg_word_length']:.1f} Avg Word Length")

                st.markdown("**Resume Preview:**")
                preview_text = row['text'][:st.session_state.preview_length]
                if len(row['text']) > st.session_state.preview_length:
                    preview_text += "..."
                st.text_area("Resume Preview", preview_text, height=150, disabled=True, key=f"preview_{i}", label_visibility="collapsed")
        else:
            # Semantic similarity details
            with st.expander(f"#{i} - {row['filename']} | Score: {row['similarity_score']:.3f} | {row['status']}"):
                col1, col2, col3 = st.columns(3)
                with col1:
                    st.metric("Similarity Score", f"{row['similarity_score']:.3f}")
                with col2:
                    st.metric("Category", row['category'])
                with col3:
                    st.markdown(
                        f"**Status:** <span style='color: {status_color}; font-weight: bold'>{row['status']}</span>",
                        unsafe_allow_html=True)

                st.markdown("**Document Metrics:**")
                # Utilize the text statistics utility function
                stats = calculate_text_statistics(row['text'])

                stat_col1, stat_col2, stat_col3 = st.columns(3)
                with stat_col1:
                    st.caption(f"📝 {stats['words']} Words")
                with stat_col2:
                    st.caption(f"📑 {stats['sentences']} Sentences")
                with stat_col3:
                    st.caption(f"📏 {stats['avg_word_length']:.1f} Avg Word Length")

                st.markdown("**Resume Preview:**")
                preview_text = row['text'][:st.session_state.preview_length]
                if len(row['text']) > st.session_state.preview_length:
                    preview_text += "..."
                st.text_area("Resume Preview", preview_text, height=150, disabled=True, key=f"preview_{i}", label_visibility="collapsed")

    # Export results
    st.header("💾 Export Results")

    if matching_type == "Hybrid Matching (Similarity + ML Classifier)":
        export_df = results_df[
            ['filename', 'category', 'similarity_score', 'classifier_confidence', 'hybrid_score', 'status']]
        csv_export = export_df.to_csv(index=False)
        filename = "hybrid_resume_screening_results.csv"
    else:
        export_df = results_df[['filename', 'category', 'similarity_score', 'status']]
        csv_export = export_df.to_csv(index=False)
        filename = "semantic_resume_screening_results.csv"

    st.download_button(
        label="📥 Download Results as CSV",
        data=csv_export,
        file_name=filename,
        mime="text/csv"
    )


def get_status_color(status: str) -> str:
    """Get color for status display"""
    return "#28a745" if status == "PASS" else "#dc3545"


def display_instructions():
    """Display usage instructions"""
    st.header("📖 How to Use")

    with st.expander("🚀 Quick Start Guide"):
        st.markdown("""
        ### Step-by-Step Instructions:

        1. **Choose Data Source:**
           - Upload individual resume files (PDF, DOCX, TXT)
           - Upload a CSV file with resume text
           - Use the pre-loaded dataset (if available)

        2. **Enter Job Description:**
           - Paste the job description directly
           - Or upload a job description file

        3. **Configure Matching:**
           - Choose between Semantic Similarity or Hybrid Matching
           - Adjust weights for hybrid matching (if available)
           - Set the pass/fail threshold

        4. **Run Analysis:**
           - Click "Match Resumes" to start the analysis
           - View results with interactive visualizations
           - Export results as CSV

        ### Features:
        - **Semantic Similarity:** Uses TF-IDF and cosine similarity
        - **Hybrid Matching:** Combines similarity with ML classification
        - **Pass/Fail Criteria:** Configurable threshold for screening
        - **Interactive Visualizations:** Charts and detailed breakdowns
        - **Export Capabilities:** Download results for further analysis
        """)

    with st.expander("🔧 Advanced Features"):
        st.markdown("""
        ### Configuration Options:

        - **TF-IDF Features:** Adjust the number of features for text vectorization
        - **N-gram Range:** Configure word combinations for analysis
        - **Display Options:** Customize number of results and preview length
        - **Threshold Tuning:** Fine-tune pass/fail criteria

        ### File Support:
        - **PDF:** Automatic text extraction from PDF documents
        - **DOCX:** Microsoft Word document processing
        - **TXT:** Plain text file support
        - **CSV:** Bulk resume data import

        ### Hybrid Matching:
        - Combines semantic similarity with ML classification
        - Weighted scoring system (α for similarity, β for classifier)
        - Category prediction for job descriptions
        - Enhanced accuracy through dual approaches
        """)

    with st.expander("⚠️ Troubleshooting"):
        st.markdown("""
        ### Common Issues:

        **File Upload Problems:**
        - Ensure files are under 10MB
        - Check file format (PDF, DOCX, TXT supported)
        - Verify text content is extractable

        **Low Similarity Scores:**
        - Review job description quality
        - Check resume text extraction
        - Consider adjusting TF-IDF parameters

        **Hybrid Matching Unavailable:**
        - Ensure ML models are properly loaded
        - Check file paths in configuration
        - Verify model compatibility

        **Performance Issues:**
        - Reduce number of resumes for large datasets
        - Lower TF-IDF max features if needed
        - Use semantic similarity for faster processing
        """)


def create_sample_data():
    """Create sample data for demonstration"""
    sample_resumes = [
        {
            "filename": "john_doe_resume.pdf",
            "text": "John Doe is a experienced software engineer with 5 years of experience in Python, machine learning, and data science. He has worked on various projects involving natural language processing, computer vision, and web development using frameworks like Django and Flask.",
            "category": "Data Science"
        },
        {
            "filename": "jane_smith_resume.pdf",
            "text": "Jane Smith is a marketing professional with expertise in digital marketing, social media management, and content creation. She has managed campaigns for Fortune 500 companies and has experience with Google Analytics, Facebook Ads, and SEO optimization.",
            "category": "Marketing"
        },
        {
            "filename": "mike_johnson_resume.pdf",
            "text": "Mike Johnson is a financial analyst with CFA certification and 7 years of experience in investment banking, portfolio management, and risk assessment. He has expertise in financial modeling, Excel VBA, and Bloomberg terminal.",
            "category": "Finance"
        }
    ]
    return pd.DataFrame(sample_resumes)


# Additional utility functions
def validate_job_description(job_description: str) -> Tuple[bool, str]:
    """Validate job description input"""
    if not job_description or not job_description.strip():
        return False, "Job description cannot be empty"

    if len(job_description.strip()) < 50:
        return False, "Job description should be at least 50 characters long"

    return True, "Valid job description"


def calculate_text_statistics(text: str) -> Dict[str, int]:
    """Calculate basic text statistics"""
    words = text.split()
    sentences = text.split('.')

    return {
        "characters": len(text),
        "words": len(words),
        "sentences": len([s for s in sentences if s.strip()]),
        "avg_word_length": sum(len(word) for word in words) / len(words) if words else 0
    }


def format_file_size(size_bytes: int) -> str:
    """Format file size in human readable format"""
    if size_bytes == 0:
        return "0 B"

    size_names = ["B", "KB", "MB", "GB"]
    import math
    i = int(math.floor(math.log(size_bytes, 1024)))
    p = math.pow(1024, i)
    s = round(size_bytes / p, 2)
    return f"{s} {size_names[i]}"


# Run the application
if __name__ == "__main__":
    main()
